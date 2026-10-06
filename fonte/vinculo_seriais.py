import json,sys,collections
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
D=json.load(open('/home/claude/work/data.json'));R=D['rows'];C=D['codes'];N=len(R)
DREN={'R':'DR','P':'DP'}
LID=['AC-C','F23-C','F1-C','F1-D','F23-D','AC-D','DR-C','DR-D'];LNAME=['Acostamento','Faixa 2/3','Faixa 1','Faixa 1','Faixa 2/3','Acostamento','Dreno','Dreno'];SIDE=['Crescente']*3+['Decrescente']*3+['Crescente','Decrescente']
def code(L,i):
    if L<6:c=R[i][6][L];return C[c-1] if c else None
    return DREN.get(R[i][7][2 if L==6 else 1])
def nm(i,extra=0):
    k,o=R[i][0].split('+');return f"{k}+{int(o)+extra:03d}"
P=[]
for L in range(8):
    i=0;seq=0;km=None
    while i<N:
        c=code(L,i)
        if not c:i+=1;continue
        j=i
        while j+1<N and code(L,j+1)==c and R[j+1][5]==R[i][5]:j+=1
        if R[i][5]!=km:km=R[i][5];seq=0
        seq+=1
        end=R[j+1][0] if j+1<N else nm(j,20)
        P.append(dict(L=L,code=c,a=R[i][3],b=R[j][3]+20,ini=R[i][0],fim=end,tag=f"{R[i][5]}-{LID[L]}-{seq:02d}",bloco=R[i][4]));i=j+1
A=json.load(open(sys.argv[1],encoding='utf-8'))
LANE={('Crescente','Faixa 1'):2,('Crescente','Faixa 3'):1,('Crescente','Faixa 2'):1,('Decrescente','Faixa 1'):3,('Decrescente','Faixa 3'):4,('Decrescente','Faixa 2'):4,('Crescente','Acostamento'):0,('Decrescente','Acostamento'):5}
num=lambda v:None if v in (None,'') else float(str(v).replace(',','.'))
REC=[]
for o in A:
    sol=o.get('Solução','');s=sol.lower()
    fam='FS' if 'superficial' in s else 'FE' if 'estrutural' in s else 'FF' if 'funcional' in s else 'DRN' if 'dreno' in s else None
    if fam=='DRN':cd={'Raso':'DR','Profundo':'DP'}.get(o.get('Tipo de Dreno'));L={'Crescente':6,'Decrescente':7}.get(o.get('Sentido'))
    else:cd=fam;L=LANE.get((o.get('Sentido'),o.get('Faixa')))
    ka,kb=num(o.get('km inicial')),num(o.get('km final'))
    r=dict(o=o,code=cd,L=L,a=None if ka is None else round(min(ka,kb)*1000),b=None if ka is None else round(max(ka,kb)*1000))
    REC.append(r)
for r in REC:
    r['panos']=[p for p in P if r['L'] is not None and p['L']==r['L'] and p['code']==r['code'] and r['a'] is not None and min(p['b'],r['b'])-max(p['a'],r['a'])>0.5]
    ps=r['panos']
    if r['L'] is None or r['code'] is None or r['a'] is None:r['sit']='Sem dados de faixa/solução/km'
    elif not ps:r['sit']='Sem pano R08 com a mesma solução'
    elif len(ps)==1:
        p=ps[0];r['sit']='Igual ao pano' if (p['a'],p['b'])==(r['a'],r['b']) else 'Dentro do pano' if p['a']<=r['a'] and r['b']<=p['b'] else 'Ultrapassa o pano'
    else:r['sit']=f'Abrange {len(ps)} panos (divididos no marco km)'
for p in P:p['recs']=[r for r in REC if p in r['panos']]
alvo=[p for p in P if p['code'] in ('FF','FE','FS','DR','DP')]
wb=Workbook();ws=wb.active;ws.title='Panos R08 x serial'
H=['Pano','Bloco','Sentido','Faixa / lado','Solução','Estaca inicial','Estaca final','Extensão (m)','Serial','Status no apontamento','Vínculo']
ws.append(H)
cnt=collections.Counter()
for p in alvo:
    rs=p['recs'];dec=SIDE[p['L']]=='Decrescente'
    v='Único (automático)' if len(rs)==1 else 'Sem apontamento' if not rs else f'{len(rs)} apontamentos: escolher na programação'
    cnt[v.split(':')[0]]+=1
    ws.append([p['tag'],'B'+str(p['bloco']),SIDE[p['L']],LNAME[p['L']],p['code'],p['fim'] if dec else p['ini'],p['ini'] if dec else p['fim'],p['b']-p['a'],' / '.join(r['o']['Serial'] for r in rs),' / '.join(str(r['o'].get('Status') or '') for r in rs),v])
ws2=wb.create_sheet('Apontamentos x R08')
ws2.append(['Serial','Solução','Tipo de dreno','Sentido','Faixa','Estaca inicial','Estaca final','km inicial','km final','Status','Identificação','Pano(s) R08','Vínculo'])
c2=collections.Counter()
for r in REC:
    o=r['o'];c2[r['sit'].split(' (')[0]]+=1
    ws2.append([o['Serial'],o.get('Solução'),o.get('Tipo de Dreno'),o.get('Sentido'),o.get('Faixa'),o.get('Estaca inicial'),o.get('Estaca final'),o.get('km inicial'),o.get('km final'),o.get('Status'),o.get('Identificação'),', '.join(p['tag'] for p in r['panos']),r['sit']])
ws3=wb.create_sheet('Resumo',0)
ws3.append(['Vínculo dos seriais (apontamentos BR-277) com os panos do R08 – Blocos 2 e 3']);ws3.append([])
ws3.append(['Panos R08 de FF, FE, FS, DR e DP',len(alvo)])
for k,v in cnt.most_common():ws3.append(['  '+k,v])
ws3.append([]);ws3.append(['Apontamentos com serial',len(REC)])
for k,v in c2.most_common():ws3.append(['  '+k,v])
ws3.append([]);ws3.append(['Critério: mesma rodovia, sentido, faixa (Faixa 3 = Faixa 2/3; drenos = lado da pista), solução (FF, FE, FS, Dreno Raso/Profundo) e trecho sobreposto pelo km (hodômetro). PA, RF, REC e RP não constam dos apontamentos.'])
for w in (ws,ws2):
    for c in w[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1A1A1A')
    w.freeze_panes='A2';w.auto_filter.ref=w.dimensions
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width=min(42,max(10,max(len(str(c.value or '')) for c in col[:300])+2))
ws3.column_dimensions['A'].width=60;ws3['A1'].font=Font(bold=True,size=13)
wb.save(sys.argv[2])
print(cnt,c2)
